"""Shared helpers for ClinixRAG ingestion and embeddings.

Heavy third-party imports are intentionally kept inside the functions so the
Flask app can boot (and show a friendly setup screen) before optional
dependencies or API keys are configured.
"""

from __future__ import annotations

from pathlib import Path
from typing import List

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx", ".epub", ".rtf"}

# Kept in one place so the app, the ingestion script and the UI agree.
CHUNK_SIZE = 800
CHUNK_OVERLAP = 100
EMBEDDING_MODEL = "sentence-transformers/all-MiniLM-L6-v2"


def load_pdf_file(data):
    """Load every PDF found in a directory (original helper, kept for parity)."""
    documents = []
    for path in sorted(Path(data).glob("*.pdf")):
        documents.extend(load_pdf_documents(path))
    return documents


def load_pdf_documents(file_path):
    """Load a single PDF, one LangChain document per page."""
    from langchain_core.documents import Document
    from pypdf import PdfReader

    path = Path(file_path)
    reader = PdfReader(str(path))
    documents = []
    for index, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception:  # pragma: no cover - malformed page
            text = ""
        if text.strip():
            documents.append(
                Document(
                    page_content=text,
                    metadata={"source": path.name, "page": index},
                )
            )
    return documents


def text_split(extracted_data, chunk_size: int = CHUNK_SIZE, chunk_overlap: int = CHUNK_OVERLAP):
    """Split LangChain documents into retrieval-sized chunks."""
    try:
        from langchain_text_splitters import RecursiveCharacterTextSplitter
    except ImportError:  # pragma: no cover - older LangChain installs
        from langchain.text_splitter import RecursiveCharacterTextSplitter

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    return splitter.split_documents(extracted_data)


def download_hugging_face_embeddings():
    """Return the sentence-transformer embeddings used by the existing index."""
    try:
        from langchain_huggingface import HuggingFaceEmbeddings
    except ImportError:  # pragma: no cover - older LangChain installs
        from langchain_community.embeddings import HuggingFaceEmbeddings

    return HuggingFaceEmbeddings(model_name=EMBEDDING_MODEL)


def _read_text_file(file_path: Path) -> str:
    raw = file_path.read_bytes()
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _load_docx(file_path: Path, filename: str):
    import docx
    from langchain_core.documents import Document

    document = docx.Document(str(file_path))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    text = "\n\n".join(paragraphs)
    return [Document(page_content=text, metadata={"source": filename})]


def _load_epub(file_path: Path, filename: str):
    import ebooklib
    from ebooklib import epub
    from langchain_core.documents import Document

    book = epub.read_epub(str(file_path))
    chapters: List[str] = []
    for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
        try:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(item.get_content(), "html.parser")
            text = soup.get_text("\n", strip=True)
        except ImportError:  # very small fallback when bs4 is unavailable
            text = item.get_content().decode("utf-8", errors="ignore")
        if text.strip():
            chapters.append(text)

    if not chapters:
        return []
    return [
        Document(page_content="\n\n".join(chapters), metadata={"source": filename})
    ]


def extract_documents(file_path, filename: str | None = None):
    """Extract text from an uploaded book into LangChain Documents.

    Supported formats: PDF, DOCX, TXT, Markdown, EPUB and RTF.
    """
    from langchain_core.documents import Document

    path = Path(file_path)
    display_name = filename or path.name
    extension = Path(display_name).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{extension or 'unknown'}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if extension == ".pdf":
        documents = load_pdf_documents(path)
    elif extension == ".docx":
        documents = _load_docx(path, display_name)
    elif extension == ".epub":
        documents = _load_epub(path, display_name)
    elif extension == ".rtf":
        # RTF is plain enough for a best-effort text extraction.
        import re

        text = _read_text_file(path)
        text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
        text = re.sub(r"[{}]", " ", text)
        documents = [Document(page_content=text, metadata={"source": display_name})]
    else:  # .txt / .md
        documents = [
            Document(
                page_content=_read_text_file(path),
                metadata={"source": display_name},
            )
        ]

    cleaned = []
    for document in documents:
        if document.page_content and document.page_content.strip():
            document.metadata.setdefault("source", display_name)
            cleaned.append(document)

    if not cleaned:
        raise ValueError("No readable text found in this file.")
    return cleaned


def human_file_size(num_bytes: int | float) -> str:
    """Format bytes as a short human readable string."""
    value = float(num_bytes or 0)
    for unit in ("B", "KB", "MB", "GB"):
        if value < 1024 or unit == "GB":
            return f"{value:.0f} {unit}" if unit == "B" else f"{value:.1f} {unit}"
        value /= 1024
    return f"{value:.1f} GB"


def allowed_file(filename: str) -> bool:
    return Path(filename).suffix.lower() in SUPPORTED_EXTENSIONS
